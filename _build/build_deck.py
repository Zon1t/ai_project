"""SR 처리 에이전트 기획서 (표지 + 4장) 생성: python build_deck.py <output.pptx>"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

FONT, MONO = "Malgun Gothic", "Consolas"
INK, MUTED, LINE, LANE, WHITE = "16201C", "5A6762", "CDD6D2", "ECF0EE", "FFFFFF"
AGENT, AGENT_SOFT = "0D6B60", "DCEFEA"
ML, ML_SOFT = "6A3D9E", "ECE3F6"
KNOW, KNOW_SOFT = "94640C", "F6EAD2"
FAIL, FAIL_SOFT = "C03F0B", "FCE6DB"
DARK, DARK2, MINT, PALE = "0F2A26", "1E4A43", "7FD1C3", "CFE3DE"

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(h):
    return RGBColor.from_string(h)


def style_run(run, size, bold, color, font, italic=False):
    f = run.font
    f.size, f.bold, f.italic = Pt(size), bold, italic
    f.color.rgb = rgb(color)
    f.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def fill_tf(tf, paras, size=11, color=INK, bold=False, align="l", anchor="t",
            font=FONT, margin=0.0, space=0, bullet=False, lsp=None):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = ANCHOR[anchor]
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ALIGN[align]
        if space:
            para.space_after = Pt(space)
        if lsp:
            para.line_spacing = lsp
        runs = [p] if isinstance(p, (str, tuple)) else p
        is_bullet = bullet
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            if o.get("bullet") is not None:
                is_bullet = o["bullet"]
            run = para.add_run()
            run.text = t
            style_run(run, o.get("size", size), o.get("bold", bold), o.get("color", color),
                      o.get("font", font), o.get("italic", False))
        if is_bullet:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.16)))
            pPr.set("indent", str(-Inches(0.14)))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")


def text(slide, x, y, w, h, paras, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, paras, **kw)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, radius=0.08, dash=False,
         paras=None, **kw):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    kw.setdefault("margin", 0.08)
    fill_tf(shp.text_frame, paras if paras is not None else "", **kw)
    return shp


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=1.25, head=True, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if tail:
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if head:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def _border(side, color):
    ln = etree.Element(qn(f"a:{side}"), w=str(Pt(0.75)) if color else "0")
    if color:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
    else:
        etree.SubElement(ln, qn("a:noFill"))
    return ln


def table(slide, x, y, colw, rows, size=9.5, hsize=9, rowh=0.27, fills=None,
          first_bold=True, mono_cols=(), header=True):
    nr, nc = len(rows), len(colw)
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)),
                                Inches(rowh * nr))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for j, w in enumerate(colw):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(rowh)
        is_head = header and i == 0
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            bg = LANE if is_head else WHITE
            if fills and (i, j) in fills:
                bg = fills[(i, j)]
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(bg)
            tf = cell.text_frame
            if is_head:
                fill_tf(tf, val, size=hsize, bold=True, color=MUTED)
            else:
                fnt = MONO if j in mono_cols else FONT
                fill_tf(tf, val, size=size, bold=(first_bold and j == 0), font=fnt)
            tf.margin_left = tf.margin_right = Inches(0.07)
            tf.margin_top = tf.margin_bottom = Inches(0.04)
            tcPr = cell._tc.get_or_add_tcPr()
            for side, col in reversed([("lnL", None), ("lnR", None), ("lnT", None),
                                       ("lnB", LINE)]):
                tcPr.insert(0, _border(side, col))
    return gf


def header(slide, eyebrow, title, page):
    text(slide, 0.5, 0.3, 12.3, 0.3, [[(eyebrow, {"bold": True, "color": AGENT, "size": 11})]])
    text(slide, 0.5, 0.58, 12.3, 0.62, title, size=24, bold=True, anchor="m")
    text(slide, 0.5, 7.1, 8, 0.25, "변경 요청 처리 에이전트 · AI Agent 프로젝트 기획서",
         size=9, color=MUTED)
    text(slide, 10.83, 7.1, 2.0, 0.25, f"{page} / 4", size=9, color=MUTED, align="r")


def section(slide, x, y, w, label, color=AGENT):
    text(slide, x, y, w, 0.28, [[(label, {"bold": True, "color": color, "size": 12})]])


def build(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---------- 표지 ----------
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(DARK)
    text(s, 0.7, 1.35, 11, 0.35, [[("SCSA27 · AI AGENT 프로젝트 기획서",
                                    {"bold": True, "color": MINT, "size": 13})]])
    text(s, 0.7, 1.75, 11.5, 1.0, "변경 요청 처리 에이전트", size=44, bold=True, color=WHITE,
         anchor="m")
    text(s, 0.7, 2.8, 11.5, 0.5,
         "결함 위치를 예측하고, 고친 코드를 스스로 공격해 검증하는 유지보수 코드 수정 에이전트",
         size=18, color=PALE)
    steps = ["분석", "수정", "공격", "판정", "승인"]
    for i, st in enumerate(steps):
        x = 0.7 + i * 1.75
        rect(s, x, 3.75, 1.35, 0.46, fill=FAIL if st == "공격" else DARK2, radius=0.5,
             paras=st, size=14, bold=True, color=WHITE, align="c", anchor="m")
        if i < len(steps) - 1:
            arrow(s, x + 1.4, 3.98, x + 1.7, 3.98, color=MINT, w=1.5)
    reqs = [("① 현업 연계", "SI·ITO 운영 조직의 변경 요청(SR) 처리"),
            ("② ML/DL", "결함 위치 예측 (DL 임베딩 + ML 순위)"),
            ("③ Tool 호출", "분석 에이전트가 예측 모델을 도구로 호출"),
            ("④ 화면", "Streamlit 티켓 처리 · 승인 화면")]
    for i, (k, v) in enumerate(reqs):
        rect(s, 0.7 + i * 3.0, 4.9, 2.85, 1.0, fill=DARK2, radius=0.08, margin=0.15,
             paras=[[(k, {"bold": True, "color": MINT, "size": 11})],
                    [(v, {"color": WHITE, "size": 12})]], space=4)
    text(s, 0.7, 6.55, 11, 0.35, "작성자  (이름 기입)      ·      2026. 10", size=13, color=PALE)

    # ---------- 1. 개요 · 문제 + 개선 · 효과 ----------
    s = prs.slides.add_slide(blank)
    header(s, "① 프로젝트 개요 · 현 문제점   +   ② 개선 방향 · 기대 효과",
           "SR 한 건의 시간은 원인 찾기와 회귀 확인에서 샌다", 1)
    L = 0.5
    section(s, L, 1.35, 4.2, "대상 업무 · 사용자")
    text(s, L, 1.66, 4.2, 0.8, [
        [("사용자  ", {"bold": True}), "운영 시스템을 유지보수하는 SI·ITO 개발자"],
        [("대상 업무  ", {"bold": True}), "변경 요청(SR) 처리"],
        [("MVP 범위  ", {"bold": True}), "결함 수정 · 노후화 대응 · 보안 진단 조치"]],
         size=10.5, space=3)
    section(s, L, 2.55, 4.2, "현재 업무 흐름")
    flow = [("1 티켓 접수", False), ("2 원인 위치 탐색", True), ("3 코드 수정", False),
            ("4 테스트·회귀 확인", True), ("5 규칙·이력 대조", True), ("6 승인·배포", False)]
    for i, (lab, pain) in enumerate(flow):
        r, c = divmod(i, 3)
        x, y = L + c * 1.45, 2.88 + r * 0.5
        rect(s, x, y, 1.27, 0.38, fill=FAIL_SOFT if pain else LANE,
             line=FAIL if pain else None, radius=0.2, paras=lab, size=9.5,
             bold=pain, color=FAIL if pain else INK, align="c", anchor="m", margin=0.03)
        if c < 2:
            arrow(s, x + 1.29, y + 0.19, x + 1.43, y + 0.19, w=1)
    section(s, L, 3.95, 4.2, "현업 문제점", color=FAIL)
    stats = [("1~2시간", ["SR당 결함 원인 위치를 찾는 데", "쓰는 시간 (로그·코드 검색)"]),
             ("회귀 결함", ["고치면서 다른 동작을 깨뜨리고,", "통합 테스트 단계에서야 발견"]),
             ("수작업", ["사내 규칙과 과거 장애 조치 이력을", "사람이 직접 찾아 대조"])]
    for i, (big, desc) in enumerate(stats):
        y = 4.28 + i * 0.78
        rect(s, L, y, 4.2, 0.68, fill=FAIL_SOFT, radius=0.1)
        text(s, L + 0.12, y, 1.45, 0.68, big, size=19, bold=True, color=FAIL, anchor="m")
        text(s, L + 1.6, y, 2.5, 0.68, desc, size=10, anchor="m")
    text(s, L, 6.66, 4.2, 0.3, "※ 문제 수치와 E2E 해결률 목표는 가정치입니다. 근거 확인 후 교체",
         size=9, color=KNOW, bold=True)

    R, RW = 4.95, 7.88
    rect(s, R, 1.35, RW, 0.66, fill=AGENT_SOFT, radius=0.12, margin=0.14, anchor="m",
         paras=[[("해결 아이디어   ", {"bold": True, "color": AGENT, "size": 11}),
                 ("모델이 고칠 위치를 좁히고, 에이전트가 이력·규칙을 근거로 수정한 뒤, "
                  "공격 에이전트가 미달·회귀를 찾을 때까지 검증합니다.",
                  {"bold": True, "size": 11.5})]])
    ba = [["단계", "Before", "After", "담당"],
          ["원인 위치", "로그·코드를 사람이 검색", "결함 위치 예측 모델이 함수 Top-3와 SHAP 근거 제시",
           [("ML / DL", {"bold": True, "color": ML})]],
          ["수정", "개인 경험에 의존", "유사 장애 이력·사내 규칙을 근거로 패치 생성",
           [("Agent + RAG", {"bold": True, "color": AGENT})]],
          ["검증", "테스트 수작업, 회귀는 늦게 발견",
           ["재현 테스트 자동 작성 +", "LLM·Hypothesis 입력으로 미달·회귀 탐색"],
           [("Agent + 판정 코드", {"bold": True, "color": AGENT})]],
          ["승인", "리뷰어가 전부 확인", "검증 결과·근거를 한 화면에서 보고 승인",
           [("사람 (HITL)", {"bold": True, "color": FAIL})]]]
    table(s, R, 2.13, [0.85, 1.95, 3.7, 1.38], ba, size=10)
    kpi = [["정량 목표", "목표치", "측정 방법"],
           ["결함 위치 Top-3 정확도", "0.80 이상", "홀드아웃 티켓, TF-IDF 베이스라인과 비교"],
           ["SR당 처리 시간", "2시간 → 15분", "시연 티켓의 접수부터 승인 요청까지 시간"],
           ["공격 에이전트 탐지율", "80% 이상", "일부러 불완전하게 만든 패치 10건 중 미달·회귀 탐지 수"],
           ["근거 인용률", "80% 이상", "패치 설명에 규칙 ID·이력 ID가 포함된 비율"],
           ["E2E 해결률", "60% 이상", "홀드아웃 20건 중 사람 개입 없이 판정 노드 통과 비율"],
           ["예측 도구 기여", "해결률 향상", "Ablation: 예측 도구 유무별 해결률·토큰·소요 시간"],
           ["회귀 누출", "1건 이하", "Ablation: 공격 단계 유무별 숨은 회귀 테스트에 걸린 패치 수"],
           ["자동 검증 통과", "3유형 모두", "시연 시나리오 3개가 판정 노드 통과"]]
    kfill = {(i, 1): AGENT_SOFT for i in range(1, len(kpi))}
    table(s, R, 4.0, [2.05, 1.35, 4.48], kpi, size=10, rowh=0.3, fills=kfill)

    # ---------- 2. 시스템 설계 1/2: 아키텍처 + 도구 ----------
    s = prs.slides.add_slide(blank)
    header(s, "③ 시스템 설계 (1/2) · 아키텍처와 도구 구성",
           "결함 위치 예측 모델은 분석 에이전트가 호출하는 도구다", 2)
    bands = [(1.35, 0.6, "화면"), (2.05, 1.4, "Agent"), (3.55, 1.42, "Tools"),
             (5.07, 1.13, "모델\n데이터")]
    for y, h, lab in bands:
        rect(s, 0.5, y, 6.5, h, fill=LANE, radius=0.06)
        text(s, 0.52, y, 0.66, h, lab.split("\n"), size=9.5, bold=True, color=MUTED,
             align="c", anchor="m")
    rect(s, 1.2, 1.43, 5.7, 0.44, fill=AGENT_SOFT, line=AGENT, radius=0.1, anchor="m",
         margin=0.04, align="c",
         paras=[[("Streamlit 화면   ", {"bold": True, "size": 10.5}),
                 ("티켓 목록 · 진행 단계 · 위치 Top-3 · 근거 · diff · 공격 로그 · 승인",
                  {"size": 9, "color": MUTED})]])
    nodes = [("AGENT", "분석 에이전트", "위치 예측 · 근거\n재현 테스트", 1.08, None),
             ("AGENT", "수정 에이전트", "규칙 반영 패치\n테스트 실행", 1.08, None),
             ("AGENT", "공격 에이전트", "수정 미달 · 회귀\n탐색", 1.08, None),
             ("CODE", "판정 노드", "유형별\n통과 조건", 0.93, None),
             ("HITL", "사람 승인", "반영 / 반려", 0.93, "h")]
    xs, x = [], 1.2
    for tag, title, sub, w, kind in nodes:
        xs.append((x, w))
        fill, ln, tc = (FAIL_SOFT, FAIL, FAIL) if kind == "h" else (
            (WHITE, MUTED, MUTED) if tag == "CODE" else (WHITE, AGENT, AGENT))
        rect(s, x, 2.5, w, 0.85, fill=fill, line=ln, lw=1, radius=0.1, margin=0.05,
             align="c", anchor="m",
             paras=[[(tag, {"bold": True, "size": 8, "color": tc, "font": MONO})],
                    [(title, {"bold": True, "size": 10.5})]] +
                   [[(ln_, {"size": 8.5, "color": MUTED})] for ln_ in sub.split("\n")])
        x += w + 0.15
    for i in range(len(xs) - 1):
        x0, w0 = xs[i]
        arrow(s, x0 + w0, 2.925, xs[i + 1][0], 2.925, w=1.25)
    cx = [a + b / 2 for a, b in xs]
    arrow(s, 1.45, 1.87, 1.45, 2.5)
    arrow(s, 6.6, 2.5, 6.6, 1.87)
    arrow(s, cx[3], 2.5, cx[3], 2.3, color=FAIL, w=1.5, head=False)
    arrow(s, cx[3], 2.3, cx[1], 2.3, color=FAIL, w=1.5, head=False)
    arrow(s, cx[1], 2.3, cx[1], 2.5, color=FAIL, w=1.5)
    text(s, cx[1], 2.07, cx[3] - cx[1], 0.2, "미달 · 회귀 발견 → 수정 재시도 (최대 3회)",
         size=8.5, bold=True, color=FAIL, align="c", anchor="m")
    groups = [(1.2, 2.05, "분석 도구", [("predict_fault_location", ML), ("search_history", KNOW),
                                       ("get_callers", INK), ("get_coverage", INK)]),
              (3.35, 1.2, "수정 도구", [("search_rules", KNOW), ("read_method", INK),
                                       ("apply_patch", INK), ("run_tests", INK)]),
              (4.65, 2.25, "공격 · 판정 도구", [("run_diff", INK), ("property_attack", INK),
                                           ("attack_memory", KNOW), ("check_deprecated", INK),
                                           ("check_security", INK)])]
    for gx, gw, gt, items in groups:
        rect(s, gx, 3.62, gw, 1.28, fill=WHITE, line=AGENT, radius=0.06, margin=0.08,
             paras=[[(gt, {"bold": True, "size": 9.5})]] +
                   [[(n, {"font": MONO, "size": 9, "color": c, "bold": c != INK})]
                    for n, c in items], lsp=0.95)
    for ax in (cx[0], 3.43, 4.7, cx[3]):
        arrow(s, ax, 3.35, ax, 3.62, w=1)
    data = [("DL + ML", "결함 위치 모델", "임베딩 +\nXGBoost Ranker", ML_SOFT, ML),
            ("CHROMA", "장애·조치 이력", "결함 주입\n데이터", KNOW_SOFT, KNOW),
            ("CHROMA", "사내 규칙 KB", "규칙 40건", KNOW_SOFT, KNOW),
            ("CHROMA", "공격 이력", "성공 공격\n재사용", KNOW_SOFT, KNOW),
            ("PYTHON", "주문·정산\n프로젝트", "원본 / 작업본", WHITE, MUTED)]
    dw = (5.7 - 4 * 0.12) / 5
    for i, (tag, title, sub, fill, ln) in enumerate(data):
        dx = 1.2 + i * (dw + 0.12)
        rect(s, dx, 5.13, dw, 1.0, fill=fill, line=ln, radius=0.08, margin=0.04,
             align="c", anchor="m",
             paras=[[(tag, {"bold": True, "size": 7.5, "color": ln, "font": MONO})]] +
                   [[(t, {"bold": True, "size": 9.5})] for t in title.split("\n")] +
                   [[(t, {"size": 8.5, "color": MUTED})] for t in sub.split("\n")])
        mx = dx + dw / 2
        arrow(s, mx, 5.13, mx, 4.9, w=1, tail=(i == 3))
    text(s, 0.5, 6.3, 6.5, 0.5, [
        [("■ ", {"color": ML}), ("ML·DL 도구    ", {}), ("■ ", {"color": KNOW}),
         ("RAG 도구    ", {}), ("■ ", {"color": INK}), ("결정적 코드    ", {}),
         ("━ ", {"color": FAIL, "bold": True}), ("수정 재시도 루프", {})],
        "모델이 Tool로 호출되는 흐름: 분석 에이전트 → predict_fault_location → 결함 위치 모델"],
        size=9, color=MUTED, space=2)

    section(s, 7.25, 1.35, 5.58, "도구(Tool) 명세")
    tools = [["도구", "입력 → 출력", "담당 업무"],
             ["predict_fault_location", "ticket_text, top_k=3 → [{method, score, reasons}]",
              ["결함 위치 예측", "(ML/DL 모델)"]],
             ["search_history", "query, top_k=3 → [{ticket, root_cause, fix_summary}]",
              "유사 장애 이력 검색"],
             ["get_callers", "method → [method]", "AST 호출부 조회"],
             ["get_coverage", "method → [{test, lines_hit}]", "통과 테스트 조회"],
             ["search_rules", "code, top_k=5 → [{rule_id, description, good_example}]",
              "사내 규칙 검색"],
             ["read_method\napply_patch", "method[, new_code] → 코드 / {applied, compile_ok, error}",
              "작업본 읽기·수정"],
             ["run_tests", "version, tests? → {passed, failed: [{test, trace}]}", "테스트 실행"],
             ["run_diff", "method, inputs → [{input, original, patched, same}]",
              ["LLM 설계 입력", "차등 실행"]],
             ["property_attack", "method, strategy_spec, max_examples=200 → [{input, ..., same}]",
              ["Hypothesis 입력", "차등 실행·최소화"]],
             ["attack_memory", "op, attack_type, record? → [{input, expected, actual}]",
              ["성공 공격", "저장·재사용"]],
             ["check_deprecated\ncheck_security", "version → [{location, api | rule_id}]",
              ["지원 종료 API", "보안 위반 검사"]]]
    rows = [tools[0]] + [["", [(r[1], {"font": MONO, "size": 8.5})], r[2]] for r in tools[1:]]
    t = table(s, 7.25, 1.66, [1.75, 2.38, 1.45], rows, size=9, rowh=0.3, first_bold=False)
    for i, r in enumerate(tools[1:], start=1):
        color = ML if "predict" in r[0] else (KNOW if r[0].startswith(("search", "attack")) else INK)
        cell = t.table.cell(i, 0)
        p = cell.text_frame.paragraphs[0]
        for run in list(p.runs):
            p._p.remove(run._r)
        names = r[0].split("\n")
        for k, nm in enumerate(names):
            para = p if k == 0 else cell.text_frame.add_paragraph()
            run = para.add_run()
            run.text = nm
            style_run(run, 8.5, True, color, MONO)

    # ---------- 3. 시스템 설계 2/2: 데이터 · 모델 · 운영 ----------
    s = prs.slides.add_slide(blank)
    header(s, "③ 시스템 설계 (2/2) · 데이터 · 모델 · Agent 운영",
           "결함을 주입해 만든 데이터로 위치 예측 모델을 학습한다", 3)
    section(s, 0.5, 1.35, 6.1, "데이터 명세 (공개 · 가상 생성만 사용)")
    dat = [["데이터", "출처 · 형식 · 건수", "주요 컬럼 · 의미"],
           ["주문·정산 프로젝트", "가상 생성 (Python 3.8 스타일, 3.12 실행) · 모듈 10개, 함수 약 100개",
            "주문·할인·결제·정산·고객 + pytest. utcnow·pytz 등 지원 종료 예고 API와 보안 약점 의도 포함"],
           ["결함 티켓 데이터셋", "mutmut + LLM 결함 주입 후 LLM이 티켓 작성 · CSV 300~500건",
            "ticket_id, title, body(증상), stacktrace(일부), fixed_method(정답), inject_type"],
           ["장애·조치 이력 KB", "티켓 데이터 학습 분할 재사용 · Chroma 약 250건",
            "ticket, root_cause, fix_summary, method"],
           ["사내 규칙 KB", "시큐어코딩 가이드·개발 표준 요약 + 가상 규칙 · JSON 40건",
            "rule_id, category, description, bad/good_example, severity"],
           ["공격 이력", "실행 중 자동 생성 · Chroma",
            "ticket_id, attack_type(미달/회귀), input, expected, actual"]]
    table(s, 0.5, 1.66, [1.3, 2.2, 2.6], dat, size=9)
    text(s, 0.5, 4.32, 6.1, 0.45, [
        [("누수 방지  ", {"bold": True, "color": KNOW}),
         "KB·프롬프트 예시는 학습 분할에서만 생성. 평가 티켓과 E2E 홀드아웃 20건은 KB·예시·과거 결함 수 피처에서 제외. 스택트레이스는 일부 티켓에만 포함"]],
         size=9, color=MUTED)

    section(s, 6.85, 1.35, 5.98, "ML/DL 모델 설계: 결함 위치 예측", color=ML)
    mdl = [["항목", "내용"],
           ["문제 유형", "순위 학습 (Learning to Rank): 티켓 1건당 후보 함수 순위"],
           ["입력 → 출력", "티켓 텍스트 + 함수 코드 → 함수별 결함 점수, 상위 3개"],
           ["DL 역할", ["사전학습 임베딩(UniXcoder 또는 bge-m3)으로",
                       "한국어 증상과 영어 코드의 의미 유사도 계산 (파인튜닝 없음)"]],
           ["ML 역할", ["XGBoost Ranker(LambdaMART)로 피처 종합:",
                       "임베딩·TF-IDF 유사도, 스택트레이스 일치, 복잡도·길이, 과거 결함 수"]],
           ["비교 실험", "① TF-IDF 단독(베이스라인)  ② 임베딩 단독  ③ 혼합형"],
           ["평가 · 목표", "Top-3 ≥ 0.80, Top-1, MRR · 티켓 단위 8:2 분할 · 스택트레이스·주입 방식별 보고"],
           ["예측 근거", "SHAP 기여도 상위 피처 3개를 reasons로 반환해 화면에 표시"],
           ["Ablation", "홀드아웃 20건: 예측 도구 유무(해결률·토큰·시간), 공격 단계 유무(회귀 누출)"]]
    mfill = {(i, 0): ML_SOFT for i in range(1, len(mdl))}
    table(s, 6.85, 1.66, [1.15, 4.83], mdl, size=9, fills=mfill)

    section(s, 0.5, 4.88, 6.1, "Agent 구조: 워크플로 그래프")
    text(s, 0.5, 5.17, 6.1, 0.42,
         "순서(분석→수정→공격→판정→승인)가 고정이라 Supervisor 대신 워크플로 그래프를 씁니다. "
         "통과 여부는 LLM이 아닌 판정 노드(코드)가 결정하고, 실패 시 실패 내용을 붙여 수정 에이전트를 다시 호출합니다.",
         size=9, color=INK)
    cond = [["수정 유형", "판정 노드 통과 조건"],
            ["결함 수정", ["재현 테스트 원본 실패 → 수정본 통과 · 기존 테스트 통과",
                          "결함 영역 밖 Hypothesis 입력 200개에서 출력 동일"]],
            ["노후화 대응", ["Hypothesis 입력 200개에서 출력 완전 동일",
                           "지원 종료 API 0건 (DeprecationWarning을 오류로 처리)"]],
            ["보안 조치", "bandit·사내 규칙 위반 0건 · 공격 입력 무력화 · 정상 입력은 원본과 출력 동일"]]
    table(s, 0.5, 5.65, [1.05, 5.05], cond, size=9)

    section(s, 6.85, 4.88, 5.98, "상태 · 메모리 / 안전장치 / 기술 스택")
    rect(s, 6.85, 5.2, 2.9, 1.08, fill=LANE, radius=0.08, margin=0.1, paras=[
        [("상태 · 메모리", {"bold": True, "size": 9.5, "color": AGENT})],
        [("티켓별 단계·예측 위치·패치·시도 횟수", {"bullet": True})],
        [("공격 이력 장기 저장, 같은 유형 우선 재사용", {"bullet": True})],
        [("에이전트 간 구조화 JSON만 전달", {"bullet": True})]], size=9, space=1)
    rect(s, 9.9, 5.2, 2.93, 1.08, fill=FAIL_SOFT, radius=0.08, margin=0.1, paras=[
        [("안전장치 (Guard)", {"bold": True, "size": 9.5, "color": FAIL})],
        [("HITL: 반영 전 담당자 승인", {"bullet": True})],
        [("테스트·기대 출력 수정 금지", {"bullet": True})],
        [("임시 작업본 · 타임아웃 10초 · 재시도 3회", {"bullet": True})]], size=9, space=1)
    rect(s, 6.85, 6.36, 5.98, 0.62, fill=WHITE, line=LINE, radius=0.08, margin=0.1, anchor="m",
         paras=[[("기술 스택  ", {"bold": True, "color": AGENT}),
                 ("Python 3.12 · Claude claude-opus-5-5 (Anthropic SDK tool use) · XGBoost · scikit-learn · "
                  "sentence-transformers · SHAP · pytest · coverage.py · Hypothesis · mutmut · bandit · ruff · "
                  "Chroma · SQLite · Streamlit", {})]], size=8.5)

    # ---------- 4. 기능 범위 · 화면 · 일정 ----------
    s = prs.slides.add_slide(blank)
    header(s, "④ 기능 범위 · 화면 · 일정",
           "필수 기능 5개로 시연 시나리오 3개를 완주한다", 4)
    section(s, 0.5, 1.35, 4.0, "기능 범위")
    feats = [["구분", "기능", "설명"],
             ["Must 1", "가상 프로젝트·실행", "주문·정산, run_tests, run_diff"],
             ["Must 2", "결함 위치 예측 모델", "결함 주입 데이터·학습·Tool화"],
             ["Must 3", "RAG 2종", "장애·조치 이력, 사내 규칙"],
             ["Must 4", "3종 에이전트 + 판정", "유형별 통과 조건, 재시도 3회"],
             ["Must 5", "Streamlit + 승인", "근거·diff·공격 로그, 승인"],
             ["Nice", "연쇄 수정", "호출부까지 다중 파일 패치"],
             ["Nice", "DL 크로스 인코더", "UniXcoder 파인튜닝과 비교"],
             ["Nice", "전략 재설계", "반복 실패 시 수정 전략 교체"],
             ["Nice", "외부 검증", "BugsInPy 일부 적용"]]
    ffill = {}
    for i in range(1, len(feats)):
        must = feats[i][0].startswith("Must")
        ffill[(i, 0)] = AGENT_SOFT if must else LANE
        feats[i][0] = [(feats[i][0], {"bold": True, "color": AGENT if must else MUTED})]
    table(s, 0.5, 1.66, [0.7, 1.45, 1.95], feats, size=9, rowh=0.27, fills=ffill)

    section(s, 4.75, 1.35, 4.05, "화면 UI 스케치 (1개 화면)")
    rect(s, 4.75, 1.66, 4.05, 2.85, fill=WHITE, line=MUTED, lw=1, dash=True, radius=0.04)
    rect(s, 4.85, 1.76, 1.2, 1.08, fill=LANE, radius=0.06, margin=0.06, paras=[
        [("티켓 목록", {"bold": True, "size": 8.5, "color": MUTED})],
        [("SR-2041 결함", {"bold": True, "color": AGENT})], "SR-2077 노후화", "SR-2103 보안"],
        size=9, space=1)
    rect(s, 4.85, 2.92, 1.2, 0.3, fill=WHITE, line=LINE, radius=0.2, paras="+ 새 티켓",
         size=9, align="c", anchor="m", margin=0.02)
    for i, (k, v) in enumerate([("재시도", "2/3"), ("미달·회귀", "1건"), ("테스트", "48/48")]):
        rect(s, 4.85, 3.3 + i * 0.38, 1.2, 0.32, fill=LANE, radius=0.1, margin=0.06, anchor="m",
             paras=[[(k + "  ", {"size": 8.5, "color": MUTED}), (v, {"bold": True, "font": MONO})]],
             size=9)
    for i, (st, kind) in enumerate([("분석", "d"), ("수정", "d"), ("공격", "n"), ("판정", ""),
                                    ("승인", "")]):
        fill = AGENT_SOFT if kind == "d" else (FAIL_SOFT if kind == "n" else WHITE)
        rect(s, 6.15 + i * 0.52, 1.76, 0.48, 0.25, fill=fill, line=LINE, radius=0.5, margin=0,
             paras=st, size=8.5, align="c", anchor="m", bold=kind == "n",
             color=FAIL if kind == "n" else INK)
    text(s, 6.15, 2.06, 2.55, 0.22, [[("결함 위치", {"bold": True, "color": AGENT}),
                                      ("  |  근거  |  diff  |  공격 로그", {"color": MUTED})]], size=8.5)
    rect(s, 6.15, 2.32, 2.55, 1.03, fill=LANE, radius=0.06, margin=0.07, paras=[
        [("결함 위치 Top-3", {"bold": True, "size": 8.5, "color": MUTED})],
        [("1 discount.apply_coupon  0.82", {"font": MONO, "bold": True, "color": ML})],
        [("2 order_calc.total_amount  0.41", {"font": MONO})],
        [("3 settlement.net_amount  0.17", {"font": MONO})],
        [("SHAP: 임베딩 +0.31 · 과거 결함 +0.12", {"size": 8, "color": MUTED})]], size=8.5)
    rect(s, 6.15, 3.42, 2.55, 0.62, fill=FAIL_SOFT, radius=0.06, margin=0.07, paras=[
        [("[미달 · Hypothesis]", {"bold": True, "color": FAIL})],
        [("amount=100000, coupons=2", {"font": MONO})],
        "기대 90,000 · 실제 -10,000"], size=8.5)
    rect(s, 6.9, 4.12, 0.6, 0.28, fill=WHITE, line=MUTED, radius=0.2, paras="반려", size=8.5,
         align="c", anchor="m", margin=0)
    rect(s, 7.58, 4.12, 1.12, 0.28, fill=AGENT, radius=0.2, paras="승인 후 반영", size=8.5,
         bold=True, color=WHITE, align="c", anchor="m", margin=0)

    section(s, 9.05, 1.35, 3.78, "5일 개발 일정 (1인)")
    sched = [["일차", "작업 → 완료 기준"],
             ["1일", [("주문·정산 프로젝트·pytest, 실행 도구, 환경(venv·mutmut) → ", {}),
                     ("run_tests·run_diff 동작", {"bold": True, "color": AGENT})]],
             ["2일", [("결함 주입 → 티켓 데이터셋, 위치 모델 학습·평가 → ", {}),
                     ("세 모델 비교표", {"bold": True, "color": AGENT})]],
             ["3일", [("RAG 2종, 분석·수정 에이전트 → ", {}),
                     ("시나리오 1 CLI 완주", {"bold": True, "color": AGENT})]],
             ["4일", [("공격 에이전트·판정, Streamlit (밤: E2E 배치 실행) → ", {}),
                     ("시나리오 1~3 화면 동작", {"bold": True, "color": AGENT})]],
             ["5일", [("오전 E2E·Ablation 정리, 오후 KPI·발표 준비 (버퍼) → ", {}),
                     ("리허설 완료", {"bold": True, "color": AGENT})]]]
    table(s, 9.05, 1.66, [0.55, 3.23], sched, size=9)

    section(s, 0.5, 4.75, 8.3, "핵심 시연 시나리오")
    scen = [("① 결함 수정", ["SR-2041 · 10만 원 이상 주문에 쿠폰 2장을", "쓰면 정산 금액이 음수"],
             ["위치 예측 → 이력 검색 → 재현 테스트", "→ 패치 → property_attack → 판정 → 승인"],
             ["Top-1 apply_coupon 예측.", "정확히 10만 원에서 1차 패치 미달 발견, 2차 통과"]),
            ("② 노후화 대응", ["SR-2077 · 3.12 전환 대비", "utcnow()·pytz 사용 제거"],
             ["check_deprecated → 규칙 검색 → 패치", "→ property_attack(월말·윤년·DST)",
              "→ 판정 → 승인"],
             "zoneinfo로 교체, 생성 입력 전체에서 원본과 출력 동일, 지원 종료 API 0건"),
            ("③ 보안 진단 조치", ["SR-2103 · find_by_name에서", "f-string으로 SQL 조립"],
             ["search_rules(SEC-SQL-01) → 바인딩 패치", "→ check_security",
              "→ run_diff(인젝션·정상) → 공격 저장"],
             "규칙 ID 인용, 인젝션 무력화, 정상 조회 동일. 성공 공격은 다음 티켓에서 재사용")]
    for i, (t1, tk, fl, rs) in enumerate(scen):
        rs = [rs] if isinstance(rs, str) else rs
        rect(s, 0.5 + i * 2.8, 5.06, 2.7, 1.93, fill=LANE, radius=0.06, margin=0.1, paras=
             [[(t1, {"bold": True, "size": 10.5, "color": AGENT})]] +
             [[(ln_, {"color": MUTED, "italic": True})] for ln_ in tk] +
             [[(ln_, {"color": AGENT, "bold": True})] for ln_ in fl] +
             [[("결과  ", {"bold": True}), (rs[0], {})]] + [[(ln_, {})] for ln_ in rs[1:]],
             size=8.5, space=2)

    section(s, 9.05, 4.75, 3.78, "리스크와 대응", color=FAIL)
    risks = [["리스크", "대응"],
             ["2일차 과부하", "mutmut 결과 위주, LLM 주입은 보조"],
             ["가상 데이터 과대평가", "주입 방식별로 나눠 보고"],
             ["패치 미수렴", "재시도 3회 후 사람 검토"],
             ["LLM 코드의 환경 침범", "작업본 격리 · 타임아웃 10초"],
             ["Hypothesis 거짓 불일치", "금액 정수 · UTC 고정"],
             ["시연 중 API 지연", "시연 티켓 결과 캐시 재생"]]
    table(s, 9.05, 5.06, [1.5, 2.28], risks, size=8.5, rowh=0.26)

    prs.save(out)


if __name__ == "__main__":
    build(sys.argv[1])
